from __future__ import annotations

import csv
import io
import math
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .exporters import BASE_RAC_HEADERS, BASE_RECORD_FIELDS
from .logic import VAV_SD_PARAMETERS
from .models import DeviceRecord, ExtraParameter
from .project_io import INTEGER_FIELDS


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    return str(value)


def _integer(value: str):
    if not value.strip():
        return None
    try:
        number = Decimal(value.strip())
    except InvalidOperation as exc:
        raise ValueError("expected a whole number") from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError("expected a whole number")
    return int(number)


def _number(value: str):
    if not value.strip():
        return None
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("expected a finite number")
    return number


def _boolean(value: str):
    normalized = value.strip().casefold()
    if not normalized:
        return None
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    raise ValueError("expected True/False, Yes/No, 1/0, or blank")


def parse_rac_rows(rows) -> list[DeviceRecord]:
    """Read SCT v3's six header rows without recalculating imported values."""
    rows = [[_text(value) for value in row] for row in rows]
    if len(rows) < 6:
        raise ValueError("Expected an SCT Rapid Archive schedule with six header rows.")
    width = max(len(row) for row in rows[:6])
    headers = [row + [""] * (width - len(row)) for row in rows[:6]]
    if not width or headers[1][0].strip().casefold() != "v3":
        raise ValueError("Only SCT Rapid Archive v3 schedules are supported. Keep the original six header rows.")

    field_ids = dict(zip(BASE_RAC_HEADERS[:25], BASE_RECORD_FIELDS))
    known_parameters = {(attribute, kind): field for _name, attribute, kind, field in VAV_SD_PARAMETERS}
    columns = {}
    parameters = {}
    used_fields = set()
    used_parameters = set()
    parameter_section = False
    for index in range(width):
        identifier = headers[0][index].strip().upper()
        if identifier in field_ids:
            field = field_ids[identifier]
            if field in used_fields:
                raise ValueError(f"Duplicate schedule column: {identifier}.")
            used_fields.add(field)
            columns[index] = field
        elif identifier == "RAC-4448" or (parameter_section and not identifier):
            parameter_section = True
            name, attribute, kind = (headers[row][index].strip() for row in (3, 4, 5))
            if not any((name, attribute, kind)):
                continue
            if not attribute or not kind:
                raise ValueError(f"Parameter column {index + 1} needs an Attribute ID and Attribute Type.")
            key = (attribute, kind)
            if key in used_parameters:
                raise ValueError(f"Duplicate parameter column: {attribute} / {kind}.")
            used_parameters.add(key)
            parameters[index] = (name, attribute, kind, known_parameters.get(key))
        elif identifier or any(headers[row][index].strip() for row in (3, 4, 5)):
            raise ValueError(f"Unsupported schedule column {index + 1}: {identifier or headers[3][index]}. No rows were imported.")
    if not {"device_name", "equipment_name"} <= used_fields or not parameter_section:
        raise ValueError("This is not an SCT Rapid Archive schedule. Device Name, Equipment Name, and RAC-4448 headers are required.")

    records = []
    for row_number, values in enumerate(rows[6:], start=7):
        if not any(value.strip() for value in values):
            continue
        if any(value.strip() for value in values[width:]):
            raise ValueError(f"Row {row_number} contains values beyond the schedule's headers.")
        values = values[:width] + [""] * max(0, width - len(values))
        for index, value in enumerate(values):
            if value.strip() and index not in columns and index not in parameters:
                raise ValueError(f"Row {row_number}, column {index + 1}: value has no recognized column header.")
        record = DeviceRecord(manufacturer="", preserve_imported_values=True)
        for index, field in columns.items():
            raw = values[index]
            try:
                value = _integer(raw) if field in INTEGER_FIELDS else _boolean(raw) if field == "dhcp_enabled" else raw
            except ValueError as exc:
                raise ValueError(f"Row {row_number}, {field}: {exc}.") from exc
            setattr(record, field, value)
        for index, (name, attribute, kind, field) in parameters.items():
            raw = values[index]
            if field is not None:
                try:
                    value = _number(raw)
                except ValueError as exc:
                    raise ValueError(f"Row {row_number}, {name}: expected a finite number or blank.") from exc
                setattr(record, field, value)
                record.parameters[name] = value
                # Retain a blank column's definition when saving an unfinished CSV.
                if value is None:
                    record.extra_parameters.append(ExtraParameter(name, attribute, kind))
            else:
                record.extra_parameters.append(ExtraParameter(name, attribute, kind, raw))
        records.append(record)
    if not records and parameters:
        raise ValueError("This schedule has parameter definitions but no device rows. Add a device row before importing it.")
    return records


def import_rac_schedule(path: str | Path) -> list[DeviceRecord]:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if "Rapid Archive Schedule" not in workbook.sheetnames:
                raise ValueError("The workbook must contain a 'Rapid Archive Schedule' sheet. Other CWD/Excel layouts are not supported.")
            rows = []
            for row in workbook["Rapid Archive Schedule"].iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("The Rapid Archive Schedule sheet contains formulas. Replace them with values or import an SCT CSV.")
                rows.append([cell.value for cell in row])
            return parse_rac_rows(rows)
        finally:
            workbook.close()
    if path.suffix.lower() != ".csv":
        raise ValueError("Choose an SCT .csv file or a RAC .xlsx workbook.")
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True))
    except (csv.Error, IndexError) as exc:
        raise ValueError("Cannot read this CSV. Expected the original SCT Rapid Archive headers and comma-separated values.") from exc
    return parse_rac_rows(rows)

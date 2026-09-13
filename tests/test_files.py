import copy
import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from rac_generator.exporters import BASE_RECORD_FIELDS, _rac_rows, export_rac_csv, export_rac_workbook
from rac_generator.importers import import_rac_schedule, parse_rac_rows
from rac_generator.logic import recalculate_record
from rac_generator.models import DeviceRecord, ExtraParameter
from rac_generator.project_io import decode_project, load_project, project_data, save_project


class FileTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.record = DeviceRecord(
            site_hierarchy="Campus/Building/Floor", room_number="001", leaf_space="Office",
            device_name="Controller-01", fqr="Existing.Custom.FQR", device_description="Café, floor\nWest",
            equipment_name="VAV-01", controller_part="Part", engine_name="SNE03", trunk_name="FC-1",
            controller_host_name="Host", mac_address=4, zigbee_pan_offset="00", instance=54321,
            n2_address="09", eth1="A", eth2="B", equipment_definition="Definition", controller_template="VAV-RH",
            manufacturer="Titus", inlet_size=8, sa_area=0.375, sa_kfactor=2.77,
            clg_maxflow=1000, clg_minflow=0, htg_minflow=100,
            mechanical_drawing="M-001", jci_ctrl_dwg_no="C-001", sensor_code_no="S-001",
            box_heat=True, supplemental_heat=False, comments="Finish tomorrow",
            parameters={"custom metadata": {"value": "keep"}},
            extra_parameters=[ExtraParameter("CUSTOM", "AV9999", "Default Value", "=literal, text\n001")],
        )

    def assert_schedule_equal(self, expected, actual):
        for field in BASE_RECORD_FIELDS + ["sa_area", "sa_kfactor", "clg_maxflow", "clg_minflow", "htg_minflow"]:
            with self.subTest(field=field):
                self.assertEqual(getattr(actual, field), getattr(expected, field))
        self.assertEqual(actual.extra_parameters, expected.extra_parameters)
        self.assertTrue(actual.preserve_imported_values)

    def test_project_preserves_complete_records_settings_and_incomplete_forms(self):
        settings = {
            "site": "", "start": "not finished", "maxflow": "1.", "manufacturer": "Titus", "inlet": "8",
            "network_type": "IP", "dhcp": False, "existing_equipment": "AHU-01\nAHU-02\n",
        }
        self.record.preserve_imported_values = True
        path = self.directory / "unfinished.rac.json"
        save_project(path, settings, [self.record, DeviceRecord()])
        restored_settings, restored_records = load_project(path)
        self.assertEqual(restored_settings, settings)
        self.assertEqual(restored_records, [self.record, DeviceRecord()])
        save_project(path, settings, [])
        self.assertEqual(load_project(path), (settings, []))

    def test_project_rejects_unsupported_or_invalid_data(self):
        valid = project_data({}, [self.record])
        bad_values = [
            ("version", 2), ("version", True), ("format", "other"), ("unknown", "cannot drop"),
            ("settings", {"unknown": "value"}), ("settings", {"dhcp": "false"}),
            ("settings", {"network_type": "unknown"}), ("records", {}),
        ]
        for key, value in bad_values:
            with self.subTest(key=key, value=value):
                invalid = copy.deepcopy(valid)
                invalid[key] = value
                with self.assertRaises(ValueError):
                    decode_project(invalid)
        for field, value in [
            ("instance", True), ("sa_area", float("nan")), ("device_name", []),
            ("extra_parameters", [{"name": "incomplete"}]), ("unknown", "cannot drop"),
        ]:
            with self.subTest(field=field):
                invalid = copy.deepcopy(valid)
                invalid["records"][0][field] = value
                with self.assertRaises(ValueError):
                    decode_project(invalid)
        path = self.directory / "invalid.rac.json"
        for contents in ['{"format": 1, "format": 2}', '{"value": NaN}', '{broken']:
            path.write_text(contents, encoding="utf-8")
            with self.assertRaises(ValueError):
                load_project(path)

    def test_failed_project_or_csv_save_preserves_previous_file(self):
        for name, save in [
            ("work.rac.json", lambda path: save_project(path, {}, [self.record])),
            ("work.csv", lambda path: export_rac_csv(path, [self.record])),
        ]:
            with self.subTest(name=name):
                path = self.directory / name
                path.write_text("previous saved work", encoding="utf-8")
                with patch("rac_generator.file_io.os.replace", side_effect=OSError("Disk unavailable")):
                    with self.assertRaises(OSError):
                        save(path)
                self.assertEqual(path.read_text(), "previous saved work")
                self.assertEqual(list(self.directory.glob("*.tmp")), [])

    def test_csv_round_trip_keeps_identifiers_network_fields_and_parameters(self):
        ip = copy.deepcopy(self.record)
        ip.equipment_name, ip.device_name = "VAV-02", "Controller-02"
        ip.mac_address, ip.ip_controller_number, ip.dhcp_enabled = None, 0, False
        ip.ip_address, ip.subnet_mask, ip.ip_router = "192.168.1.10", "255.255.255.0", "192.168.1.1"
        path = self.directory / "working.csv"
        export_rac_csv(path, [self.record, ip, DeviceRecord(device_name="Unfinished")])
        loaded = import_rac_schedule(path)
        self.assertEqual(len(loaded), 3)
        for expected, actual in zip([self.record, ip], loaded):
            self.assert_schedule_equal(expected, actual)
        self.assertEqual(loaded[2].equipment_name, "")
        self.assertEqual(loaded[0].manufacturer, "")
        self.assertIsNone(loaded[0].inlet_size)
        project = self.directory / "resumed.rac.json"
        save_project(project, {}, loaded)
        export_rac_csv(path, load_project(project)[1])
        self.assertEqual(import_rac_schedule(path), loaded)

    def test_empty_working_csv_and_blank_parameter_columns(self):
        path = self.directory / "empty.csv"
        export_rac_csv(path, [])
        self.assertEqual(import_rac_schedule(path), [])
        blank = DeviceRecord(device_name="Incomplete", extra_parameters=[
            ExtraParameter("SA-AREA", "AV3111", "Default Value"),
            ExtraParameter("CUSTOM", "AV9999", "Default Value"),
        ])
        export_rac_csv(path, [blank])
        imported = import_rac_schedule(path)
        self.assertEqual(imported[0].extra_parameters, blank.extra_parameters)
        export_rac_csv(path, imported)
        self.assertEqual(import_rac_schedule(path), imported)

    def test_csv_encodings_delimiters_and_reordered_base_columns(self):
        rows = _rac_rows([self.record])
        for row in rows:
            row[1], row[3] = row[3], row[1]
        for delimiter, encoding in [(";", "cp1252"), ("\t", "utf-16"), (",", "utf-8-sig")]:
            with self.subTest(delimiter=delimiter, encoding=encoding):
                output = io.StringIO(newline="")
                csv.writer(output, delimiter=delimiter).writerows(rows)
                path = self.directory / "working.csv"
                path.write_bytes(output.getvalue().encode(encoding))
                self.assert_schedule_equal(self.record, import_rac_schedule(path)[0])

    def test_invalid_schedule_data_is_rejected_instead_of_dropped(self):
        mutations = [
            (1, 0, "v4"), (0, 1, "RAC-9999"), (0, 1, "RAC-4096"),
            (6, 12, "4.5"), (6, 17, "maybe"), (6, 25, "nan"),
            (4, 25, ""), (4, 26, "AV3111"),
        ]
        for row, col, value in mutations:
            with self.subTest(row=row, col=col, value=value):
                rows = _rac_rows([self.record])
                rows[row][col] = value
                with self.assertRaises(ValueError):
                    parse_rac_rows(rows)
        for rows in [[], [[] for _ in range(6)], _rac_rows([self.record])[:6]]:
            with self.assertRaises(ValueError):
                parse_rac_rows(rows)
        rows = _rac_rows([self.record])
        rows[6].append("unmapped data")
        with self.assertRaises(ValueError):
            parse_rac_rows(rows)

    def test_imported_values_survive_recalculation_and_unrelated_edits(self):
        record = parse_rac_rows(_rac_rows([self.record]))[0]
        data = {"Titus": {8: (0.35, 2.39)}}
        recalculate_record(record, data, False, "Custom workbook convention")
        self.assert_schedule_equal(self.record, record)
        record.room_number = "002"
        recalculate_record(record, data, changed_field="room_number")
        self.assertEqual((record.fqr, record.instance, record.sa_area, record.sa_kfactor),
                         ("Existing.Custom.FQR", 54321, 0.375, 2.77))
        self.assertIn("002", record.device_description)
        record.manufacturer, record.inlet_size = "Titus", 8
        recalculate_record(record, data, changed_field="inlet_size")
        self.assertEqual((record.sa_area, record.sa_kfactor), (0.35, 2.39))
        self.assertEqual(record.extra_parameters, self.record.extra_parameters)

    def test_master_workbook_round_trip_retains_extra_parameters_and_literal_text(self):
        self.record.device_description = "=literal text"
        self.record.extra_parameters.extend(
            ExtraParameter(f"=CUSTOM-{number}", f"AV{9000 + number}", "Default Value", str(number))
            for number in range(12)
        )
        path = self.directory / "master.xlsx"
        template = Path(__file__).resolve().parents[1] / "resources" / "RAC Schedule_Template_Updated.xlsx"
        export_rac_workbook(path, template, [self.record])
        self.assert_schedule_equal(self.record, import_rac_schedule(path)[0])
        workbook = load_workbook(path)
        try:
            self.assertEqual(workbook["Generated Scratchpad"].cell(2, 6).data_type, "s")
        finally:
            workbook.close()

    def test_unrelated_workbooks_and_formulas_are_rejected(self):
        path = self.directory / "other.xlsx"
        workbook = Workbook()
        workbook.save(path)
        with self.assertRaisesRegex(ValueError, "Rapid Archive Schedule"):
            import_rac_schedule(path)
        sheet = workbook.active
        sheet.title = "Rapid Archive Schedule"
        for row in _rac_rows([self.record]):
            sheet.append(row)
        sheet.cell(7, 6).value = "=1+1"
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "formulas"):
            import_rac_schedule(path)


if __name__ == "__main__":
    unittest.main()
